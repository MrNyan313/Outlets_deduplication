"""Excel file reading and writing handler using openpyxl and xlsxwriter."""

from pathlib import Path
from typing import List, Tuple, Dict
import openpyxl
import xlsxwriter


def find_input_excel_file(root_dir: Path) -> Path:
    """
    Finds the single .xlsx file in the root project directory.
    Ignores temporary Excel files (starting with '~$') and hidden files.
    """
    candidates = [
        f for f in root_dir.glob("*.xlsx")
        if not f.name.startswith("~$") and not f.name.startswith(".")
    ]

    if not candidates:
        raise FileNotFoundError(f"No .xlsx file found in project root: {root_dir}")

    if len(candidates) > 1:
        names = [f.name for f in candidates]
        raise ValueError(
            f"Expected exactly 1 Excel file in project root, but found {len(candidates)}: {names}"
        )

    return candidates[0]


def normalize_header(header: str) -> str:
    """Replaces narrow non-breaking spaces and strips whitespace."""
    if not header:
        return ""
    return str(header).replace('\u202f', ' ').replace('\u00a0', ' ').strip()


def read_excel_data(file_path: Path) -> Tuple[List[str], List[tuple], Dict[str, int]]:
    """
    Reads the Excel file in streaming read-only mode.
    Returns:
      - raw_headers: original headers list
      - rows: list of raw row tuples
      - col_indices: dictionary mapping normalized header -> column index
    """
    wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    sheet = wb.active

    row_iterator = sheet.iter_rows(values_only=True)
    header_tuple = next(row_iterator, None)
    if not header_tuple:
        raise ValueError(f"Excel file is empty: {file_path}")

    raw_headers = [str(h) if h is not None else "" for h in header_tuple]

    # Map normalized header names to index
    col_indices: Dict[str, int] = {}
    for idx, h in enumerate(raw_headers):
        norm_h = normalize_header(h)
        col_indices[norm_h] = idx
        # Also store original if different
        if h != norm_h:
            col_indices[h] = idx

    # Required columns check
    required = ['id', 'Name', 'AccountingSystemAddress', 'IsVendor', 'Дистрибьютор']
    missing = [req for req in required if req not in col_indices]
    if missing:
        raise KeyError(f"Missing required columns in Excel: {missing}. Available: {list(col_indices.keys())}")

    rows: List[tuple] = []
    for row in row_iterator:
        rows.append(row)

    wb.close()
    return raw_headers, rows, col_indices


def write_excel_results(
    output_path: Path,
    original_headers: List[str],
    results: List[Tuple[int, str, tuple]]
) -> None:
    """
    Writes deduplication results using xlsxwriter in constant_memory mode.
    Prepends 'Итоговый ID' and 'Статус' as the first two columns.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = xlsxwriter.Workbook(str(output_path), {'constant_memory': True})
    worksheet = workbook.add_worksheet("Deduplication_Results")

    # Header style
    header_format = workbook.add_format({
        'bold': True,
        'bg_color': '#D9E1F2',
        'border': 1
    })

    full_headers = ['Итоговый ID', 'Статус'] + list(original_headers)
    for col_num, header_name in enumerate(full_headers):
        worksheet.write(0, col_num, header_name, header_format)

    # Write data rows
    for row_num, (final_id, status, original_row) in enumerate(results, start=1):
        worksheet.write(row_num, 0, final_id)
        worksheet.write(row_num, 1, status)
        for col_num, cell_val in enumerate(original_row, start=2):
            if cell_val is None:
                worksheet.write_blank(row_num, col_num, "")
            elif isinstance(cell_val, (int, float)):
                worksheet.write_number(row_num, col_num, cell_val)
            else:
                worksheet.write_string(row_num, col_num, str(cell_val))

    workbook.close()
